#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка калькулятора: HTML из каркаса + книга в корне."""
from __future__ import annotations

import json
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
ARHIV = ROOT / "_архив"
DATA_PATH = HERE / "blocks_data.json"
TPL_PATH = HERE / "каркас.html"
HTML_OUT = HERE / "calc.html"
XLSX_OUT = ROOT / "Калькулятор_содержания_ребёнка.xlsx"

NAVY = "1F3864"
BLUE = "4472C4"
HEAD2 = "D9E2F3"
YELLOW = "FFF2CC"
GREEN = "E2EFDA"
GRAY = "F2F2F2"
INK = "1A1A1A"
MUTED = "808080"
LINE = "D0D7DE"
BRAND = "21A038"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
fill_title = PatternFill("solid", fgColor=HEAD2)
fill_head = PatternFill("solid", fgColor=BLUE)
fill_sec = PatternFill("solid", fgColor=NAVY)
fill_in = PatternFill("solid", fgColor=YELLOW)
fill_calc = PatternFill("solid", fgColor=GREEN)
fill_const = PatternFill("solid", fgColor=GRAY)
fill_sample = PatternFill("solid", fgColor="21A038")

font_h1 = Font(name="Arial", size=15, bold=True, color=NAVY)
font_sub = Font(name="Arial", size=9, color=MUTED)
font_sec = Font(name="Arial", size=11, bold=True, color="FFFFFF")
font_th = Font(name="Arial", size=10, bold=True, color="FFFFFF")
font_th2 = Font(name="Arial", size=10, bold=True, color=NAVY)
font_id = Font(name="Consolas", size=10)
font_n = Font(name="Arial", size=10)
font_b = Font(name="Arial", size=10, bold=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")
money_fmt = '# ##0" ₽"'
num_fmt = "# ##0.00"


PALETTE = [
    ("Нейтральные · текст и линии", None, None, None),
    ("--c-ink", "#1A1A1A", "1A1A1A", "основной текст"),
    ("--c-ink2", "#374151", "374151", "заголовки таблиц"),
    ("--c-ink3", "#4B5563", "4B5563", "вторичный текст"),
    ("--c-gr", "#6B7280", "6B7280", "подписи"),
    ("--c-gr2", "#9CA3AF", "9CA3AF", "мелкие подписи"),
    ("--c-gr3", "#B6BCC4", "B6BCC4", "самый светлый текст"),
    ("--c-ln", "#E5E7EB", "E5E7EB", "линии и границы"),
    ("--c-ln2", "#F0F2F5", "F0F2F5", "светлые разделители"),
    ("Нейтральные · фоны", None, None, None),
    ("--c-bg", "#F5F6F8", "F5F6F8", "фон страницы"),
    ("--c-surf", "#FAFBFC", "FAFBFC", "фон подложек"),
    ("--c-white", "#FFFFFF", "FFFFFF", "фон карточек"),
    ("Зелёный · бренд Счётикс", None, None, None),
    ("--c-g-50", "#F5FBF7", "F5FBF7", "самый светлый зелёный фон"),
    ("--c-g-100", "#EAF6EE", "EAF6EE", "фон зелёных плашек"),
    ("--c-g-200", "#D8EEDF", "D8EEDF", "рамки"),
    ("--c-g-300", "#A9DDB8", "A9DDB8", "сегменты"),
    ("--c-g-350", "#6FD183", "6FD183", "градиент полоски"),
    ("--c-g-400", "#4FC26A", "4FC26A", "градиенты"),
    ("--c-g-500", "#21A038", "21A038", "основной зелёный бренда"),
    ("--c-g-600", "#1B9331", "1B9331", "кнопки"),
    ("--c-g-700", "#158026", "158026", "акцентные цифры"),
    ("--c-g-800", "#0B6B22", "0B6B22", "тёмный градиент hero"),
    ("--c-g-900", "#14532D", "14532D", "текст на зелёном"),
    ("Красный · удаление, сброс", None, None, None),
    ("--c-r-50", "#FFF6F6", "FFF6F6", "фон"),
    ("--c-r-500", "#EF4444", "EF4444", "ошибка"),
    ("--c-r-700", "#B91C1C", "B91C1C", "кнопка сброса"),
    ("Янтарный · ввод в книге", None, None, None),
    ("--c-a-50", "#FFFBEF", "FFFBEF", "фон"),
    ("--c-a-100", "#FEF3D6", "FEF3D6", "жёлтое = ввод (Excel)"),
    ("Синий · вторые полоски", None, None, None),
    ("--c-b-400", "#5CA8E0", "5CA8E0", "полоска группы"),
    ("--c-b-500", "#3B82F6", "3B82F6", "акцент"),
]

TEXTS = [
    ("h1_01", "форма", "Сколько стоит содержать ребёнка", "заголовок"),
    ("bt_01", "форма", "счетикс", "слово знака"),
    ("bt_02", "форма", "твоя юнит-экономика", "подпись знака"),
    ("sub_01", "форма", "Калькулятор считает прямые траты на ребёнка и его долю в общих расходах семьи. Цифры можно менять — итог пересчитается сразу.", "лид"),
    ("rlb_01", "форма", "Полное содержание одного ребёнка", "подпись hero"),
    ("h2_01", "форма", "Куда уходят деньги", "карточка 01"),
    ("hint_01", "форма", "Расходы на ребёнка в месяц. Нажмите «+», чтобы раскрыть группу, или название — чтобы перейти к строкам.", "подсказка диаграммы"),
    ("h2_02", "форма", "Как мы считаем", "карточка 02"),
    ("how_01", "форма", "В расчёте учитываются все расходы на содержание ребёнка — и те, что идут только на него, и его доля в общих тратах семьи.", "ввод методики"),
]


def load_data() -> dict:
    return json.loads(DATA_PATH.read_text(encoding="utf-8"))


def row_year(r: dict) -> float:
    price = float(r.get("price") or 0)
    pcs = float(r.get("pcs") or 1) or 1
    years = float(r.get("years") or 0)
    split = float(r.get("split") or 1) or 1
    if years:
        return price * pcs / years / split
    period = r.get("period") or ""
    if period == "неделя":
        return price * pcs * 52 / split
    if period == "месяц":
        return price * pcs * 12 / split
    if period == "год":
        return price * pcs / split
    return 0.0


def next_prev(folder: Path, stem: str, suffix: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    nums = []
    for p in folder.glob(f"{stem}_prev*{suffix}"):
        m = re.search(r"prev(\d+)", p.name)
        if m:
            nums.append(int(m.group(1)))
    n = (max(nums) if nums else 0) + 1
    return folder / f"{stem}_prev{n}{suffix}"


def archive_file(path: Path, folder: Path, stem: str) -> None:
    if not path.exists():
        return
    folder.mkdir(parents=True, exist_ok=True)
    dest = next_prev(folder, stem, path.suffix)
    shutil.copy2(path, dest)
    print("archive", dest.relative_to(ROOT))


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def title_block(ws, title, sub, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws["A1"] = title
    ws["A1"].font = font_h1
    ws["A2"] = sub
    ws["A2"].font = font_sub
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 32


def section_row(ws, row, text, last_col):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row, 1, text)
    cell.font = font_sec
    cell.fill = fill_sec
    for c in range(1, last_col + 1):
        ws.cell(row, c).fill = fill_sec
        ws.cell(row, c).border = thin
    ws.row_dimensions[row].height = 22


def header_row(ws, row, headers):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h)
        cell.font = font_th
        cell.fill = fill_head
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[row].height = 26
    ws.freeze_panes = f"A{row + 1}"
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"


def paint(cell, kind="n"):
    cell.border = thin
    cell.alignment = left
    if kind == "id":
        cell.font = font_id
    elif kind == "b":
        cell.font = font_b
    else:
        cell.font = font_n


def input_cell(cell, value, money=False):
    cell.value = value
    cell.fill = fill_in
    cell.border = thin
    cell.font = font_n
    cell.alignment = center
    if money:
        cell.number_format = money_fmt


def calc_cell(cell, value, money=False):
    cell.value = value
    cell.fill = fill_calc
    cell.border = thin
    cell.font = font_n
    cell.alignment = right if money else center
    if money:
        cell.number_format = money_fmt


def const_cell(cell, value, money=False):
    cell.value = value
    cell.fill = fill_const
    cell.border = thin
    cell.font = font_n
    cell.alignment = center
    if money:
        cell.number_format = money_fmt


def page(ws, landscape=False):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.horizontalCentered = True
    ws.oddFooter.right.text = "Страница &P из &N · Счётикс"
    ws.sheet_view.showGridLines = False


def validate(data: dict) -> None:
    blocks = data["blocks"]
    ids = {b["id"] for b in blocks}
    for p in data.get("parts") or []:
        for bid in p.get("blocks") or []:
            if bid not in ids:
                raise SystemExit(f"часть {p['id']}: нет блока {bid}")
    for g in data.get("chart_groups") or []:
        for bid in g.get("blocks") or []:
            if bid not in ids:
                raise SystemExit(f"группа {g['id']}: нет блока {bid}")
    if "__DATA__" not in TPL_PATH.read_text(encoding="utf-8"):
        raise SystemExit("в каркасе нет метки __DATA__")


def build_html(data: dict) -> None:
    archive_file(HTML_OUT, ARHIV / "вёрстка", "calc")
    payload = {
        "blocks": data["blocks"],
        "tiles": data.get("tiles") or [],
        "chart_groups": data.get("chart_groups") or [],
        "parts": data.get("parts") or [],
    }
    tpl = TPL_PATH.read_text(encoding="utf-8")
    html = tpl.replace("__DATA__", json.dumps(payload, ensure_ascii=False))
    HTML_OUT.write_text(html, encoding="utf-8")
    print("html", HTML_OUT.relative_to(ROOT))


def build_excel(data: dict) -> None:
    archive_file(XLSX_OUT, ARHIV / "excel", "Калькулятор_содержания_ребёнка")
    blocks = data["blocks"]
    parts = data.get("parts") or []
    groups = data.get("chart_groups") or []
    by_id = {b["id"]: b for b in blocks}
    part_of = {}
    for p in parts:
        for bid in p.get("blocks") or []:
            part_of[bid] = p

    wb = Workbook()

    # ---------- 00_Читать ----------
    ws = wb.active
    ws.title = "00_Читать"
    title_block(ws, "СЧЁТИКС · калькулятор содержания ребёнка",
                "Сколько стоит содержать одного ребёнка в семье. Обновлено 13 августа 2026.", 2)
    section_row(ws, 4, "ИЗ ЧЕГО СОСТОИТ ПРОДУКТ — три файла", 2)
    header_row(ws, 5, ["Файл", "Что это и где живёт"])
    files = [
        ("Калькулятор_содержания_ребёнка.xlsx", "ЭТА КНИГА. Логика, параметры, каталог, методика. Жёлтое можно менять."),
        ("web/calc.html", "КАЛЬКУЛЯТОР. Автономный HTML, стили и данные вшиты. Собирается, руками не править."),
        ("web/каркас.html + blocks_data.json", "ИСТОЧНИК. Вёрстка и цифры. После правки: python3 web/собрать.py"),
    ]
    for i, (a, b) in enumerate(files, 6):
        paint(ws.cell(i, 1, a), "b")
        paint(ws.cell(i, 2, b))
        ws.row_dimensions[i].height = 32
    section_row(ws, 10, "КАК УСТРОЕНА КНИГА — 10 листов", 2)
    header_row(ws, 11, ["Лист", "Назначение"])
    # freeze already set by header_row to A12 — but we have two tables.
    # reset freeze to A4 like photographer read sheet
    sheets_help = [
        ("00_Читать", "Этот лист — карта проекта"),
        ("01_Глоссарий", "Термины: амортизация, доля, минимум, период"),
        ("02_Параметры", "Члены семьи и константы. ID в колонке A. FamilySize — именованный диапазон"),
        ("03_Каталоги", "Все позиции одной таблицей. Без строк «Итого» — итог на 05_Сводка через SUMIF"),
        ("04_Блоки", "Блоки, разделы, упаковки, привязка к частям экрана"),
        ("05_Сводка", "Суммы по блокам и частям. Формулы SUMIF по каталогу"),
        ("06_Методика", "Как из цены получается год и месяц. Цепочка расчёта"),
        ("07_Тексты", "Надписи калькулятора. Колонка «Где» — поверхность"),
        ("08_Палитра", "Цвета формы. Те же токены, что у калькулятора фотографа"),
        ("09_Интерфейс", "Какие блоки на экране и в каком порядке"),
    ]
    for i, (a, b) in enumerate(sheets_help, 12):
        paint(ws.cell(i, 1, a), "id")
        paint(ws.cell(i, 2, b))
        ws.row_dimensions[i].height = 28
    section_row(ws, 23, "ЦВЕТА ЯЧЕЕК", 2)
    header_row(ws, 24, ["Цвет", "Значение"])
    ws.cell(25, 1, "Жёлтое")
    ws.cell(25, 1).fill = fill_in
    paint(ws.cell(25, 1), "b")
    paint(ws.cell(25, 2, "Ввод. Цену, количество, срок, период можно менять."))
    ws.cell(26, 1, "Зелёное")
    ws.cell(26, 1).fill = fill_calc
    paint(ws.cell(26, 1), "b")
    paint(ws.cell(26, 2, "Считается формулой. Не затирать."))
    ws.cell(27, 1, "Серое")
    ws.cell(27, 1).fill = fill_const
    paint(ws.cell(27, 1), "b")
    paint(ws.cell(27, 2, "Константа или справочное. Обычно не трогают."))
    widths(ws, {"A": 42, "B": 92})
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = None
    page(ws)

    # ---------- 01_Глоссарий ----------
    ws = wb.create_sheet("01_Глоссарий")
    title_block(ws, "01 · ГЛОССАРИЙ — понятия расчёта",
                "Одни и те же слова на экране, в книге и в методике.", 4)
    section_row(ws, 4, "A · МАРКЕРЫ ТИПОВ", 4)
    header_row(ws, 5, ["Наименование", "ID", "Тип", "Описание"])
    glossary_a = [
        ("Константа", "const", "meta", "Аксиома. Пользователь не меняет. 12 месяцев, 52 недели."),
        ("Значение по умолчанию", "default", "meta", "Предзаполнено, можно изменить. Состав семьи, цены."),
        ("Формула", "calc", "meta", "Вычисляется из других величин. Ввод недоступен."),
        ("Переключатель", "switch", "meta", "Список: неделя / месяц / год. Меняет ветку расчёта."),
    ]
    r = 6
    for name, gid, typ, desc in glossary_a:
        paint(ws.cell(r, 1, name), "b")
        paint(ws.cell(r, 2, gid), "id")
        paint(ws.cell(r, 3, typ), "id")
        paint(ws.cell(r, 4, desc))
        r += 1
    section_row(ws, r, "B · ДЕНЬГИ И ВРЕМЯ", 4)
    r += 1
    header_row(ws, r, ["Наименование", "ID", "Тип", "Описание"])
    r += 1
    glossary_b = [
        ("Срок службы", "years", "default", "Сколько лет вещь служит. Если заполнено — период не используется."),
        ("Период", "period", "switch", "неделя / месяц / год. Сколько раз в год платят, если нет срока службы."),
        ("Амортизация", "amort", "calc", "Цена × количество ÷ срок службы. Раскладываем покупку на годы."),
        ("Платёж", "pay", "calc", "Цена × количество × (52 | 12 | 1) в зависимости от периода."),
        ("Семья в год", "family_year", "calc", "Полная сумма на всю семью за год, до деления."),
        ("Доля ребёнка", "child_year", "calc", "Семья в год ÷ членов семьи. Для индивидуальных split = 1."),
        ("В месяц", "child_month", "calc", "Год ÷ 12. Чтобы сравнивать разовое и регулярное."),
        ("Минимум", "direct", "calc", "Только индивидуальные траты. Без жилья, машины, общего быта."),
        ("Полное содержание", "grand", "calc", "Минимум плюс доля в общих расходах семьи."),
        ("Членов семьи", "family_size", "default", "На сколько делим общее. По умолчанию 3."),
        ("Упаковка", "pack", "meta", "direct / housing / transport / household / trips. Минимум = только direct."),
    ]
    for name, gid, typ, desc in glossary_b:
        paint(ws.cell(r, 1, name), "b")
        paint(ws.cell(r, 2, gid), "id")
        paint(ws.cell(r, 3, typ), "id")
        paint(ws.cell(r, 4, desc))
        ws.row_dimensions[r].height = 28
        r += 1
    widths(ws, {"A": 28, "B": 16, "C": 12, "D": 88})
    ws.freeze_panes = "A6"
    page(ws)

    # ---------- 02_Параметры ----------
    ws = wb.create_sheet("02_Параметры")
    title_block(ws, "02 · ПАРАМЕТРЫ — каждая величина объявлена один раз",
                "Жёлтое — ввод. Зелёное — считается. Серое — константа. Формулы пишем через ID.", 7)
    section_row(ws, 4, "1 · СЕМЬЯ", 7)
    header_row(ws, 5, ["ID", "Название", "Значение", "Ед.", "Формула / как получено", "Тип", "Комментарий"])
    const_cell(ws.cell(6, 1), "family_size")
    ws.cell(6, 1).font = font_id
    paint(ws.cell(6, 2, "Членов семьи"))
    input_cell(ws.cell(6, 3), 3)
    paint(ws.cell(6, 4, "чел"))
    paint(ws.cell(6, 5, "на сколько делим общее"))
    paint(ws.cell(6, 6, "default"), "id")
    paint(ws.cell(6, 7, "Именованный диапазон FamilySize. Каталог ссылается на него."))
    section_row(ws, 8, "2 · КАЛЕНДАРЬ", 7)
    ws.cell(9, 1, "ID")
    for c, h in enumerate(["ID", "Название", "Значение", "Ед.", "Формула / как получено", "Тип", "Комментарий"], 1):
        cell = ws.cell(9, c, h)
        cell.font = font_th2
        cell.fill = fill_title
        cell.alignment = center
        cell.border = thin
    params = [
        ("days_year", "Дней в году", 365, "дн", "", "const", "календарный год"),
        ("months", "Месяцев в году", 12, "мес", "", "const", "перевод года в месяц"),
        ("weeks_year", "Недель в году", 52, "нед", "округлено, не 365/7", "const", "еженедельные покупки"),
        ("currency", "Валюта", "RUB", "", "", "const", "все суммы в рублях"),
    ]
    for i, row in enumerate(params, 10):
        const_cell(ws.cell(i, 1), row[0])
        ws.cell(i, 1).font = font_id
        paint(ws.cell(i, 2, row[1]))
        const_cell(ws.cell(i, 3), row[2])
        paint(ws.cell(i, 4, row[3]))
        paint(ws.cell(i, 5, row[4]))
        paint(ws.cell(i, 6, row[5]), "id")
        paint(ws.cell(i, 7, row[6]))
    widths(ws, {"A": 16, "B": 28, "C": 14, "D": 8, "E": 36, "F": 10, "G": 56})
    ws.freeze_panes = "A6"
    page(ws)
    wb.defined_names.add(DefinedName(name="FamilySize", attr_text="'02_Параметры'!$C$6"))
    wb.defined_names.add(DefinedName(name="WeeksYear", attr_text="'02_Параметры'!$C$12"))
    wb.defined_names.add(DefinedName(name="MonthsYear", attr_text="'02_Параметры'!$C$11"))

    # ---------- 03_Каталоги ----------
    ws = wb.create_sheet("03_Каталоги")
    title_block(ws, "03 · КАТАЛОГИ — все позиции одной таблицей",
                "Без строк «Итого». Итог — на 05_Сводка формулой SUMIF. Жёлтое меняйте. Период: неделя / месяц / год.", 13)
    headers = [
        "Блок", "Наименование", "Кол-во", "Цена, ₽", "Срок, лет", "Период",
        "Делим на", "Ссылка", "Тип", "Платежей/год", "Семья в год, ₽",
        "Ребёнок в год, ₽", "Ребёнок в месяц, ₽",
    ]
    header_row(ws, 4, headers)
    dv = DataValidation(type="list", formula1='"неделя,месяц,год"', allow_blank=True)
    ws.add_data_validation(dv)
    cat_first = 5
    r = 5
    for b in blocks:
        shared = str(b.get("section") or "").startswith("shared")
        kind = b.get("add") or "pay"
        for item in b.get("items") or []:
            paint(ws.cell(r, 1, b["id"]), "id")
            paint(ws.cell(r, 2, item.get("name") or ""))
            input_cell(ws.cell(r, 3), item.get("pcs") if item.get("pcs") is not None else 1)
            input_cell(ws.cell(r, 4), item.get("price") or 0, money=True)
            years = item.get("years") or 0
            if years:
                input_cell(ws.cell(r, 5), years)
            else:
                input_cell(ws.cell(r, 5), None)
            period = item.get("period") or ""
            if years:
                input_cell(ws.cell(r, 6), "")
            else:
                input_cell(ws.cell(r, 6), period or "месяц")
            if shared:
                calc_cell(ws.cell(r, 7), "=FamilySize")
            else:
                const_cell(ws.cell(r, 7), item.get("split") or 1)
            url = item.get("url") or ""
            cell_u = ws.cell(r, 8, url)
            paint(cell_u)
            if url.startswith("http"):
                cell_u.hyperlink = url
                cell_u.font = Font(name="Arial", size=10, color="0563C1", underline="single")
            paint(ws.cell(r, 9, kind), "id")
            calc_cell(ws.cell(r, 10), f'=IF(F{r}="год",1,IF(F{r}="месяц",MonthsYear,IF(F{r}="неделя",WeeksYear,0)))')
            calc_cell(ws.cell(r, 11), f"=IF(N(E{r})>0,C{r}*D{r}/E{r},C{r}*D{r}*J{r})", money=True)
            calc_cell(ws.cell(r, 12), f"=IF(G{r}=0,0,K{r}/G{r})", money=True)
            calc_cell(ws.cell(r, 13), f"=IF(L{r}=0,0,L{r}/MonthsYear)", money=True)
            dv.add(ws.cell(r, 6))
            ws.row_dimensions[r].height = 18
            r += 1
    cat_last = r - 1
    ws.auto_filter.ref = f"A4:M{cat_last}"
    widths(ws, {
        "A": 20, "B": 42, "C": 10, "D": 14, "E": 12, "F": 12,
        "G": 12, "H": 36, "I": 10, "J": 14, "K": 16, "L": 18, "M": 20,
    })
    page(ws, landscape=True)
    ws.print_title_rows = "1:4"

    # ---------- 04_Блоки ----------
    ws = wb.create_sheet("04_Блоки")
    title_block(ws, "04 · БЛОКИ — состав экрана",
                "Один блок = одна таблица в калькуляторе. pack=direct входит в минимум.", 8)
    header_row(ws, 4, ["ID", "Название", "Раздел", "Упаковка", "Часть", "Школа", "Ссылки", "Тип строк"])
    for i, b in enumerate(blocks, 5):
        p = part_of.get(b["id"])
        paint(ws.cell(i, 1, b["id"]), "id")
        paint(ws.cell(i, 2, b.get("title") or ""))
        paint(ws.cell(i, 3, b.get("section") or ""), "id")
        paint(ws.cell(i, 4, b.get("pack") or "direct"), "id")
        paint(ws.cell(i, 5, p["id"] if p else ""), "id")
        paint(ws.cell(i, 6, "да" if b.get("school") else ""))
        paint(ws.cell(i, 7, "да" if b.get("links") else ""))
        paint(ws.cell(i, 8, b.get("add") or "pay"), "id")
    last_b = 4 + len(blocks)
    ws.auto_filter.ref = f"A4:H{last_b}"
    widths(ws, {"A": 22, "B": 42, "C": 16, "D": 14, "E": 14, "F": 10, "G": 10, "H": 12})
    page(ws)

    # ---------- 05_Сводка ----------
    ws = wb.create_sheet("05_Сводка")
    title_block(ws, "05 · СВОДКА — один ребёнок",
                "Все суммы — формулы SUMIF по 03_Каталоги. Правите жёлтое в каталоге, сводка обновится.", 5)
    section_row(ws, 4, "БЛОКИ", 5)
    header_row(ws, 5, ["Блок", "ID", "Упаковка", "Ребёнок в год, ₽", "Ребёнок в месяц, ₽"])
    block_first = 6
    for i, b in enumerate(blocks, 6):
        paint(ws.cell(i, 1, b.get("title") or ""))
        paint(ws.cell(i, 2, b["id"]), "id")
        paint(ws.cell(i, 3, b.get("pack") or "direct"), "id")
        calc_cell(ws.cell(i, 4), f"=SUMIF('03_Каталоги'!$A$5:$A${cat_last},B{i},'03_Каталоги'!$L$5:$L${cat_last})", money=True)
        calc_cell(ws.cell(i, 5), f"=IF(D{i}=0,0,D{i}/MonthsYear)", money=True)
    block_last = 5 + len(blocks)
    tot = block_last + 1
    ws.cell(tot, 1, "ПОЛНОЕ СОДЕРЖАНИЕ")
    ws.merge_cells(start_row=tot, start_column=1, end_row=tot, end_column=3)
    for c in range(1, 6):
        ws.cell(tot, c).fill = fill_sec
        ws.cell(tot, c).font = font_sec
        ws.cell(tot, c).border = thin
    ws.cell(tot, 4, f"=SUM(D{block_first}:D{block_last})")
    ws.cell(tot, 5, f"=SUM(E{block_first}:E{block_last})")
    ws.cell(tot, 4).number_format = money_fmt
    ws.cell(tot, 5).number_format = money_fmt
    ws.cell(tot, 4).font = font_sec
    ws.cell(tot, 5).font = font_sec

    pr = tot + 2
    section_row(ws, pr, "ЧАСТИ ЭКРАНА", 5)
    ph = pr + 1
    for c, h in enumerate(["Часть", "ID", "Состав", "Ребёнок в год, ₽", "Ребёнок в месяц, ₽"], 1):
        cell = ws.cell(ph, c, h)
        cell.font = font_th
        cell.fill = fill_head
        cell.alignment = center
        cell.border = thin
    pf = ph + 1
    for i, p in enumerate(parts):
        rr = pf + i
        paint(ws.cell(rr, 1, p.get("title") or ""), "b")
        paint(ws.cell(rr, 2, p["id"]), "id")
        paint(ws.cell(rr, 3, "; ".join(p.get("blocks") or [])))
        ids = p.get("blocks") or []
        if ids:
            bits = "+".join(
                f"SUMIF('03_Каталоги'!$A$5:$A${cat_last},\"{bid}\",'03_Каталоги'!$L$5:$L${cat_last})"
                for bid in ids
            )
            calc_cell(ws.cell(rr, 4), f"={bits}", money=True)
        else:
            calc_cell(ws.cell(rr, 4), 0, money=True)
        calc_cell(ws.cell(rr, 5), f"=IF(D{rr}=0,0,D{rr}/MonthsYear)", money=True)

    # minimum / school checks
    chk = pf + len(parts) + 1
    section_row(ws, chk, "ПРОВЕРКИ", 5)
    ws.cell(chk + 1, 1, "Минимум (pack = direct)")
    paint(ws.cell(chk + 1, 1), "b")
    calc_cell(ws.cell(chk + 1, 4),
              f"=SUMIF(C{block_first}:C{block_last},\"direct\",D{block_first}:D{block_last})", money=True)
    calc_cell(ws.cell(chk + 1, 5), f"=D{chk+1}/MonthsYear", money=True)
    school_ids = [b["id"] for b in blocks if b.get("school")]
    ws.cell(chk + 2, 1, "Школа за год (флаг school, без кружков)")
    paint(ws.cell(chk + 2, 1), "b")
    if school_ids:
        bits = "+".join(
            f"SUMIF('03_Каталоги'!$A$5:$A${cat_last},\"{bid}\",'03_Каталоги'!$L$5:$L${cat_last})"
            for bid in school_ids
        )
        calc_cell(ws.cell(chk + 2, 4), f"={bits}", money=True)
    calc_cell(ws.cell(chk + 2, 5), f"=D{chk+2}/MonthsYear", money=True)

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Блоки, в месяц"
    data_ref = Reference(ws, min_col=5, min_row=5, max_row=block_last)
    cats = Reference(ws, min_col=1, min_row=6, max_row=block_last)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    chart.height = 14
    chart.width = 18
    ws.add_chart(chart, "G5")

    widths(ws, {"A": 44, "B": 22, "C": 14, "D": 20, "E": 22})
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:E{block_last}"
    page(ws, landscape=True)

    # ---------- 06_Методика ----------
    ws = wb.create_sheet("06_Методика")
    title_block(ws, "06 · МЕТОДИКА — как из цены получается месяц",
                "Одна цепочка. Те же правила в HTML и в формулах каталога.", 3)
    section_row(ws, 4, "ШАГИ", 3)
    header_row(ws, 5, ["Шаг", "Что делаем", "Формула"])
    steps = [
        ("1. Покупка или платёж", "Смотрим, заполнен ли срок службы в годах.", "years > 0 ?"),
        ("2a. Амортизация", "Вещь служит несколько лет — раскладываем стоимость.", "цена × кол-во ÷ лет"),
        ("2b. Платёж", "Нет срока — берём период.", "неделя × 52, месяц × 12, год × 1"),
        ("3. Семья в год", "Пока без деления. Это колонка K каталога.", "family_year"),
        ("4. Доля ребёнка", "Индивидуальное не делим. Общее — на FamilySize.", "family_year ÷ split"),
        ("5. В месяц", "Чтобы сравнивать разовое и регулярное.", "год ÷ 12"),
        ("6. Минимум", "Складываем блоки с pack = direct.", "SUMIF pack=direct"),
        ("7. Полное", "Все блоки, включая жильё, транспорт, быт, поездки.", "SUM всех child_year"),
        ("8. Школа", "Только блоки с флагом school. Без кружков и секций.", "school_ward + stuff + reg + add"),
    ]
    for i, (a, b, c) in enumerate(steps, 6):
        paint(ws.cell(i, 1, a), "b")
        paint(ws.cell(i, 2, b))
        paint(ws.cell(i, 3, c), "id")
        ws.row_dimensions[i].height = 28
    section_row(ws, 16, "ЧТО НЕ ВХОДИТ", 3)
    paint(ws.cell(17, 1, "Чужое"))
    ws.merge_cells("B17:C17")
    paint(ws.cell(17, 2, "Вещи, продукты и одежда только для других членов семьи."))
    paint(ws.cell(18, 1, "Долги"))
    ws.merge_cells("B18:C18")
    paint(ws.cell(18, 2, "Платежи по кредитам в расчёт не берутся."))
    widths(ws, {"A": 28, "B": 72, "C": 44})
    page(ws)

    # ---------- 07_Тексты ----------
    ws = wb.create_sheet("07_Тексты")
    title_block(ws, "07 · ТЕКСТЫ — надписи калькулятора",
                "Колонка «Где» — поверхность. Есть автофильтр.", 4)
    header_row(ws, 4, ["ID", "Где", "Текст", "Комментарий"])
    for i, (tid, where, text, note) in enumerate(TEXTS, 5):
        paint(ws.cell(i, 1, tid), "id")
        paint(ws.cell(i, 2, where))
        paint(ws.cell(i, 3, text))
        paint(ws.cell(i, 4, note))
        ws.row_dimensions[i].height = 28
    last_t = 4 + len(TEXTS)
    ws.auto_filter.ref = f"A4:D{last_t}"
    widths(ws, {"A": 12, "B": 12, "C": 88, "D": 28})
    page(ws)

    # ---------- 08_Палитра ----------
    ws = wb.create_sheet("08_Палитра")
    title_block(ws, "08 · ПАЛИТРА — единая со Счётикс",
                "Те же токены, что у калькулятора фотографа. Менять цвет — здесь и в каркасе одновременно.", 4)
    header_row(ws, 4, ["Токен CSS", "HEX", "Образец", "Где применяется"])
    r = 5
    for token, hexv, rgb, where in PALETTE:
        if hexv is None:
            section_row(ws, r, token, 4)
            r += 1
            continue
        paint(ws.cell(r, 1, token), "id")
        paint(ws.cell(r, 2, hexv), "id")
        sample = ws.cell(r, 3, "")
        sample.fill = PatternFill("solid", fgColor=rgb)
        sample.border = thin
        paint(ws.cell(r, 4, where))
        r += 1
    widths(ws, {"A": 16, "B": 12, "C": 14, "D": 40})
    page(ws)

    # ---------- 09_Интерфейс ----------
    ws = wb.create_sheet("09_Интерфейс")
    title_block(ws, "09 · ИНТЕРФЕЙС — что на экране и в каком порядке",
                "Карточки сверху вниз. 01 и 02 — всегда открыты. Части 03+ сворачиваются.", 4)
    section_row(ws, 4, "КАРТОЧКИ", 4)
    header_row(ws, 5, ["№", "Карточка", "Содержание", "Поведение"])
    ui = [
        ("01", "Куда уходят деньги", "Полоски по группам диаграммы", "плюс раскрывает состав группы"),
        ("02", "Как мы считаем", "Методика, членов семьи, печать и сброс", "спойлеры + параметр FamilySize"),
    ]
    for i, row in enumerate(ui, 6):
        paint(ws.cell(i, 1, row[0]), "id")
        paint(ws.cell(i, 2, row[1]), "b")
        paint(ws.cell(i, 3, row[2]))
        paint(ws.cell(i, 4, row[3]))
    r = 8
    section_row(ws, r, "ЧАСТИ — с 03", 4)
    r += 1
    for c, h in enumerate(["№", "Часть", "Блоки", "Заметка"], 1):
        cell = ws.cell(r, c, h)
        cell.font = font_th
        cell.fill = fill_head
        cell.alignment = center
        cell.border = thin
    r += 1
    for i, p in enumerate(parts):
        paint(ws.cell(r, 1, f"{i+3:02d}"), "id")
        paint(ws.cell(r, 2, p.get("title") or ""), "b")
        paint(ws.cell(r, 3, ", ".join(p.get("blocks") or [])))
        paint(ws.cell(r, 4, p.get("desc") or ("доля ÷ FamilySize" if p.get("shared") else "")))
        ws.row_dimensions[r].height = 28
        r += 1
    r += 1
    section_row(ws, r, "ГРУППЫ ДИАГРАММЫ", 4)
    r += 1
    for c, h in enumerate(["ID", "Подпись", "Блоки", ""], 1):
        cell = ws.cell(r, c, h)
        cell.font = font_th
        cell.fill = fill_head
        cell.alignment = center
        cell.border = thin
    r += 1
    for g in groups:
        paint(ws.cell(r, 1, g.get("id") or ""), "id")
        paint(ws.cell(r, 2, g.get("name") or ""), "b")
        paint(ws.cell(r, 3, ", ".join(g.get("blocks") or [])))
        r += 1
    widths(ws, {"A": 10, "B": 36, "C": 72, "D": 36})
    page(ws, landscape=True)

    # order: read first already
    XLSX_OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_OUT)
    print("xlsx", XLSX_OUT.relative_to(ROOT), "sheets", wb.sheetnames)


def report(data: dict) -> None:
    grand = direct = school = 0.0
    for b in data["blocks"]:
        y = sum(row_year(r) for r in b.get("items") or [])
        grand += y
        if (b.get("pack") or "direct") == "direct":
            direct += y
        if b.get("school") or b.get("id") == "school_must":
            school += y
    print(f"минимум/мес  {direct/12:,.0f}".replace(",", " "))
    print(f"полное/мес   {grand/12:,.0f}".replace(",", " "))
    print(f"школа/год    {school:,.0f}".replace(",", " "))


def main() -> None:
    data = load_data()
    validate(data)
    build_html(data)
    build_excel(data)
    report(data)


if __name__ == "__main__":
    main()
